#!/usr/bin/env python3
"""Generate the consolidated illustrated technical report from measured JSON.

Developer dependencies: matplotlib, numpy, markdown, python-docx (optional).
No inference or benchmarks are run. Historical evidence is snapshotted so this
generator also works inside the standalone delivery package.
"""
import hashlib
import json
from pathlib import Path
import shutil
import subprocess

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
import markdown

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'reports/技术报告'
FIG = OUT/'插图'
DATA = OUT/'数据'
for p in [OUT, FIG, DATA]:
    p.mkdir(parents=True, exist_ok=True)
TITLE = 'Jetson视频工作台_部署与提速技术报告'
LEGACY = ROOT.parent/'开发归档/早期部署工程'
plt.rcParams.update({'font.family': 'Noto Sans CJK JP', 'axes.unicode_minus': False,
                     'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False})
BLUE, GREEN, ORANGE, RED = '#2683b5', '#159982', '#e7a03c', '#ce6458'
catalog = []


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def read(relative):
    p = ROOT/relative
    catalog.append(dict(path=relative, bytes=p.stat().st_size, sha256=sha(p)))
    return json.loads(p.read_text())


snapshot = DATA/'历史后端与ReID原始摘要.json'
if LEGACY.is_dir():
    historical = {}
    for name in ['uav_torch_640', 'uav_trt_640', 'coco_torch_640', 'coco_trt_640',
                 'uav_trt_reid_e2e', 'reid_comparison', 'mot20_01_trt_reid_comparison',
                 'mot20_08_trt_reid_comparison']:
        leaf = 'comparison.json' if 'comparison' in name else 'summary.json'
        p = LEGACY/f'deployment/results/{name}/{leaf}'
        historical[name] = dict(source=str(p.relative_to(ROOT.parent)),
                                sha256=sha(p), data=json.loads(p.read_text()))
    snapshot.write_text(json.dumps(historical, ensure_ascii=False, indent=2)+'\n')
historical = json.loads(snapshot.read_text())
h = {k: v['data'] for k, v in historical.items()}
c = read('reports/validation/ceiling_comparison.json')
p = read('reports/validation/performance_comparison.json')
t = read('reports/validation/tracking_ceiling.json')
q = read('reports/benchmarks/paired_benchmark.json')
lat = read('reports/benchmarks/alternating_latency.json')
cleanup = read('reports/validation/目录整理_20261015.json')
audit32 = read('reports/audit_torch_640.json')
web = {x['state']['profile']: x for x in c['web_mot20']}
order = ['quality', 'balanced', 'fast_motion', 'fast_static', 'speed']
labels = ['1280 小目标', '640 标准', '640×384 移动', '640×384 固定', '512×320 极速']
fps = lambda k: web[k]['pipeline']['processing_fps']
ap = lambda k, metric='map50': web[k]['state']['evaluation'][metric]*100


def figure(name, fig):
    fig.tight_layout()
    fig.savefig(FIG/name, dpi=160, bbox_inches='tight', facecolor='white')
    plt.close(fig)


# 1. Responsibilities and the real pipeline. Diagram is schematic, not timing.
fig, ax = plt.subplots(figsize=(11.8, 3.2)); ax.set(xlim=(0,12), ylim=(0,3)); ax.axis('off')
steps = [('读取 / 预处理', 'CPU · 下一帧准备', BLUE), ('YOLOv12 + TensorRT', 'GPU · 框 / 类别 / 分数', GREEN),
         ('BoT-SORT', 'CPU · 运动预测 / ID关联', BLUE), ('画框 / 保存 / 预览', 'CPU + NVENC硬件编码', GREEN)]
for i,(title, detail, color) in enumerate(steps):
    x = i*3+.08
    ax.add_patch(FancyBboxPatch((x,1.3),2.68,1.05,boxstyle='round,pad=.06',facecolor=color,edgecolor='none'))
    ax.text(x+1.34,1.97,title,ha='center',va='center',color='white',fontsize=12,weight='bold')
    ax.text(x+1.34,1.56,detail,ha='center',va='center',color='white',fontsize=9)
    if i<3:ax.annotate('',(x+2.99,1.83),(x+2.72,1.83),arrowprops={'arrowstyle':'->','color':'#526474','lw':2})
ax.text(6,.63,'有界队列让各阶段重叠执行；每帧检测、每帧跟踪、每帧原尺寸保存',ha='center',fontsize=12)
ax.text(6,.13,'人工 GT → 独立检测 AP 评估；ReID 为可选外观网络，当前网页关闭',ha='center',fontsize=10,color='#526474')
figure('01_系统链路.png',fig)

fig, axes = plt.subplots(1,2,figsize=(11.4,3.6))
for ax,metric,title in [(axes[0],'end_to_end_fps','早期整段流程 FPS（越大越好）'),
                       (axes[1],'infer','检测推理 ms/帧（越小越好）')]:
    for j,(backend,label,color) in enumerate([('torch','PyTorch FP32',BLUE),('trt','TensorRT FP16',GREEN)]):
        vals = [h[f'{mode}_{backend}_640'][metric] if metric!='infer' else h[f'{mode}_{backend}_640']['stages_mean_ms']['infer'] for mode in ['uav','coco']]
        bars=ax.bar(np.arange(2)+(j-.5)*.34,vals,.33,label=label,color=color)
        ax.bar_label(bars,fmt='%.2f',padding=3,fontsize=9)
    ax.set_xticks([0,1],['无人机 · 750帧','COCO例片 · 329帧']);ax.set_title(title);ax.set_ylim(0,ax.get_ylim()[1]*1.2);ax.legend(fontsize=8);ax.grid(axis='y',alpha=.15);ax.set_axisbelow(True)
figure('02_TensorRT早期对照.png',fig)

fig,ax=plt.subplots(figsize=(10.7,3.4)); x=np.arange(3)
for j,(key,label,color) in enumerate([('before','原串行实现',ORANGE),('after','有界并行 + NVENC',GREEN)]):
    bars=ax.bar(x+(j-.5)*.34,[r[key]['processing_fps'] for r in p['paired_runs']],.33,label=label,color=color)
    ax.bar_label(bars,fmt='%.2f',padding=3)
ax.set_xticks(x,[r['label'] for r in p['paired_runs']]);ax.set_ylabel('处理 FPS');ax.set_ylim(0,51);ax.legend();ax.grid(axis='y',alpha=.15);ax.set_axisbelow(True)
figure('03_相同配置流水线对照.png',fig)

fig,axes=plt.subplots(1,2,figsize=(11.5,3.4))
experiments={r['name']:r['pipeline'] for r in c['experiments']}
graphvals=[experiments['rect640_static_final']['processing_fps'],experiments['rect640_graph']['processing_fps']]
bars=axes[0].bar(['Graph 关闭','Graph 开启'],graphvals,color=[BLUE,GREEN]);axes[0].bar_label(bars,fmt='%.2f',padding=3);axes[0].set_ylim(0,125);axes[0].set_ylabel('完整流水线 FPS');axes[0].set_title('640×384 固定 · 同形状 Graph 对照')
before={r['state']['profile']:r for r in c['web_before_tuning'] if r['state'].get('source_path','').endswith('MOT20-01-raw.mp4')}
xx=np.arange(3)
for j,(kind,color) in enumerate([('before',BLUE),('after',GREEN)]):
    vv=[before[k]['pipeline']['processing_fps'] if kind=='before' else fps(k) for k in ['balanced','fast_static','speed']]
    bars=axes[1].bar(xx+(j-.5)*.34,vv,.33,label='网页调度前' if kind=='before' else '网页调度后',color=color);axes[1].bar_label(bars,fmt='%.1f',padding=3,fontsize=8)
axes[1].set_xticks(xx,['640标准','640×384固定','512×320']);axes[1].set_ylim(0,145);axes[1].set_title('真实网页：原生JSON + GIL调度 + 预览节奏');axes[1].legend(fontsize=8)
figure('04_Graph与网页调度.png',fig)

fig,ax=plt.subplots(figsize=(10.5,3.7))
for k,label,color in zip(order,labels,[BLUE,GREEN,ORANGE,RED,'#8b75ba']):
    ax.scatter(fps(k),ap(k),s=115,color=color,zorder=3)
    offset=(-8,-23) if k=='fast_motion' else (7,7)
    ax.annotate(label,(fps(k),ap(k)),xytext=offset,textcoords='offset points',fontsize=10,
                ha='right' if k=='fast_motion' else 'left')
ax.axvline(100,color='#8995a0',ls='--',lw=1);ax.text(101,64.8,'100 FPS',color='#687886');ax.set(xlim=(0,148),ylim=(50,67),xlabel='完整处理 FPS（不含加载、预热和最后AP）',ylabel='检测 AP50（%）',title='同一 MOT20-01 全部429帧：速度与精度的实际取舍');ax.grid(alpha=.18)
figure('05_速度精度取舍.png',fig)

fig,axes=plt.subplots(1,2,figsize=(11.3,3.4))
dl=c['dla']['latency']['runs'];gpu=np.mean([r['mean_ms'] for r in dl if r['label'].startswith('gpu')]);dla=next(r['mean_ms'] for r in dl if r['label']=='dla_gpu')
v=[gpu,dla];bars=axes[0].bar(['GPU FP16','DLA0 + GPU回退'],v,color=[GREEN,RED]);axes[0].bar_label(bars,fmt='%.2f ms',padding=3);axes[0].set_ylim(0,49);axes[0].set_title('DLA隔离执行：30次预热 / 100次测量')
lv=[np.mean([r['mean_ms'] for r in lat['runs'] if r['precision']==k]) for k in ['fp16','int8']]
bars=axes[1].bar(['FP16','INT8混合'],lv,color=[GREEN,RED]);axes[1].bar_label(bars,fmt='%.2f ms',padding=3);axes[1].set_ylim(0,13);axes[1].set_title('INT8交替复测：每轮60次预热 / 250次测量')
for ax in axes:ax.set_ylabel('ms/次（越小越好）');ax.grid(axis='y',alpha=.15);ax.set_axisbelow(True)
figure('06_DLA与INT8.png',fig)

fig,ax=plt.subplots(figsize=(10.4,3.2)); xx=np.arange(3)
reids=[h[k]['results'] for k in ['reid_comparison','mot20_01_trt_reid_comparison','mot20_08_trt_reid_comparison']]
for j,(label,color) in enumerate([('ReID关',GREEN),('ReID开 · PyTorch FP32',RED)]):
    bars=ax.bar(xx+(j-.5)*.34,[rs[j]['tracker_mean_ms'] for rs in reids],.33,label=label,color=color);ax.bar_label(bars,fmt='%.1f',padding=3)
ax.set_xticks(xx,['无人机 · 100帧','MOT20-01 · 429帧','MOT20-08 · 806帧']);ax.set_ylim(0,1200);ax.set_ylabel('跟踪阶段 ms/帧');ax.set_title('完全相同的检测缓存；不含检测 / 解码 / 绘图');ax.legend()
figure('07_ReID代价.png',fig)

for source,dest in [(ROOT/'reports/figures/小目标同帧对照.png',FIG/'08_小目标同帧对照.png'),
                    (LEGACY/'部署性能报告/插图/08_ReID交汇放大.png',FIG/'09_ReID交汇放大.png')]:
    if source.exists():shutil.copy2(source,dest)


def table(headers, rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |',
                      *['| '+' | '.join(map(str,r))+' |' for r in rows]])


pages=[]
def page(title, body):
    pages.append('## '+title+'\n\n'+body.strip()+'\n')


page('01　先看结论：百帧来自整条流程的改造', f'''
在 Orin NX 16GB 上，同一段 MOT20-01 的全部 429 帧，网页预览、BoT-SORT、日志和原尺寸 H.264 保存全部开启，640×384 固定机位方案实测 **{fps('fast_static'):.2f} FPS**，512×320 极速方案 **{fps('speed'):.2f} FPS**。MOT17-09 极速方案为 121.74 FPS。每帧都检测，没有跳帧。

![系统链路](插图/01_系统链路.png)

| 需求 | 实测结果 | 代价与建议 |
| --- | --- | --- |
| 保留标准640检测与GMC | 56.00 FPS / AP50 57.12% | 保留原跟踪策略 |
| 固定机位、争取百帧 | 104.17 FPS / AP50 56.33% | 矩形输入，关闭相机运动补偿 |
| 更高吞吐 | 116.75 FPS / AP50 52.96% | 更小输入，远处目标更易漏检 |
| 小目标优先 | 18.32 FPS / AP50 63.69% | 1280输入，计算量更大 |

这些数字是**处理吞吐**。把最后的 AP、加载和预热也算进去，104.17 FPS 那次任务共用 7.29 秒，即 58.85 帧/秒的整任务平均。输出仍按源视频 25 FPS 播放，两个结果视频看起来差不多，并不表示处理速度相同。

已经证明的收益分三类：TensorRT FP16 加快网络；并行、原生后处理、硬件编码和网页调度降低工程开销；矩形输入与轻量/关闭 GMC 则用部分质量换速度。INT8、DLA 和配套 ReID 在本机没有带来适合默认启用的收益。

报告日期：2026-10-15。本文汇总已有实测，本次整理报告没有重新运行整套性能实验。所有百分比、样本与测时口径在后文注明。
''')

page('02　设备、素材和计分方式', '''
设备为 **Jetson Orin NX 16GB**，Ubuntu 22.04 / L4T 36.4.7，CUDA 12.6、TensorRT 10.7、Python 3.10。最新百帧测试使用 40W、jetson_clocks 锁定最高频率，桌面应用仍在运行。早期后端、INT8 与第一轮流水线测试为 40W 动态频率；不能把跨日期的差值全部归因于某一项代码。

| 素材 | 本报告实际使用 | 能说明什么 |
| --- | --- | --- |
| MOT20-01 | 全429帧，960×540官方缩放预览，25FPS | 最新速度/AP主对照，person单类 |
| MOT17-09 | 全525帧，极速模式 | 另一个百帧网页实例 |
| VisDrone uav0000137 | 全233帧，2688×1512，30FPS | 高分辨率、密集小目标负载 |
| 无人机 MultiUAV-003 | 全750帧，640网络输入 | UAV单类链路；缺少完整GT，不给mAP |
| 早期 COCO palace例片 | 全329帧 | PyTorch / TensorRT后端对照 |

视频库共14段，含MOT17、MOT20、VisDrone和UAV；其中12段配有可用GT。MOT素材是便于选择和演示的预览片，不能称为全分辨率官方测试集成绩。

**三个时间口径必须分开。** ①隔离推理延迟：网络执行与指定拷贝/同步，不含跟踪保存。②最新处理FPS：帧数÷处理墙钟时间，含解码、检测、跟踪、绘图、日志、全部编码、队列排空与MP4索引，排除加载/预热及最后AP。③完整任务秒数：包含模型准备和最后评估，最接近点击到完成的等待时间。早期串行实现还把最终转码放在处理计时之外，图表按历史定义保留，完整任务耗时另列。

**检测AP**使用人工框，COCOeval bbox AP，101个召回点，IoU=.50:.05:.95，每帧每类最多100个框；评估置信度0.001、NMS IoU .7、全局max_det 300。AP50只要求IoU≥.5；AP50:95要求更精确的边界。它们都不是“所有框中正确框的比例”。MOT只评person；VisDrone只评兼容COCO类别子集。**IDF1**另评跨帧身份；本报告不是官方榜单提交。

实测多为单次完整视频，不能当作所有场景的保证值。目标数量、源分辨率、摄像机运动、温度和后台负载都会改变速度。
''')

torchrows=[]
for mode,label in [('uav','无人机 / 750帧'),('coco','COCO例片 / 329帧')]:
    a,b=h[f'{mode}_torch_640'],h[f'{mode}_trt_640']
    torchrows.append([label,f"{a['end_to_end_fps']:.2f} → {b['end_to_end_fps']:.2f}",f"+{(b['end_to_end_fps']/a['end_to_end_fps']-1)*100:.1f}%",f"{a['stages_mean_ms']['infer']:.2f} → {b['stages_mean_ms']['infer']:.2f}"])
page('03　TensorRT + FP16：网络快了约三倍', '''
PyTorch逐次调用网络算子；TensorRT在本机构建引擎时选择执行算法、安排内存并融合适合的计算。FP16降低部分计算和数据搬运成本。本次同时换后端并允许FP16，测到的是两者的合并收益。

本地权重 → FP32 ONNX → 本机TensorRT engine。早期两种检测器均为batch=1、640×640，允许内部FP16、保留必要FP32回退；引擎输入输出仍是FP32。开启FP16不意味着每一层都强制半精度。

![TensorRT早期对照](插图/02_TensorRT早期对照.png)

'''+table(['同素材、同参数','完整流程FPS','吞吐变化','检测推理 ms/帧'],torchrows)+'''

无人机检测由54.28降到17.61 ms，约3.08倍；COCO由58.56降到21.16 ms，约2.77倍。整体却只提升42.2%和36.4%，因为读帧、GMC、关联、绘图和编码仍在花时间。这正是后续继续提速的依据。

这组后端对照均保存视频、开启GMC、关闭ReID、排除加载预热。它不等于最新网页的56或104 FPS：代码、视频和频率条件已有变化。更换Jetson或TensorRT版本后需要本机构建并验证引擎，不能仅拷贝一个engine就假设兼容。
''')

pairedrows=[]
for r in p['paired_runs']:
    a,b=r['before'],r['after'];pairedrows.append([r['label'],f"{a['processing_fps']:.2f} → {b['processing_fps']:.2f}",f"{b['processing_fps']/a['processing_fps']:.2f}倍",f"{a['wall_seconds']:.2f} → {b['wall_seconds']:.2f}"])
page('04　为什么曾经只有2～4 FPS，怎样解决', '''
高分辨率VisDrone原串行实现每帧平均约：读帧11.68 ms、预处理15.44 ms、推理与传输90.02 ms、后处理30.41 ms、跟踪70.81 ms、绘图写视频205.05 ms；再加循环开销，整体只有2.36 FPS。GPU网络并不是唯一瓶颈，单是绘图写出就比推理更慢。

改为解码/TRT生产线程、BoT-SORT关联和绘图/日志/编码输出线程，用长度2的队列限制积压。下一帧准备、当前帧关联和上一帧保存可以重叠；GMC角点准备也提前进行。输出由原先MJPEG中间AVI再转码，改为NVENC直接H.264 MP4；失败可回退软件编码。

![同配置优化](插图/03_相同配置流水线对照.png)

'''+table(['完整视频 / 输入','处理FPS：前→后','倍率','完整任务秒：前→后'],pairedrows)+'''

三个成对实验的 `detections.jsonl`、`evaluation_detections.jsonl`、`tracks.csv` 均逐字节一致。这里没有缩小输入或跳帧，说明工程实现本身有很大优化空间。改成并行后，各线程耗时互相重叠，不能再把阶段平均毫秒直接相加换算FPS。

mAP匹配循环也改用原生double计算，遵守原COCOeval规则，不降低评估要求。MOT640完整评估由44.46降到2.71秒；这缩短点击后的等待，不能算成检测器本身变快。表中历史“处理FPS”的最终转码边界不同，完整任务秒数把这部分也计入，便于实际比较。
''')

page('05　继续压缩CPU开销、传输与GPU调度', '''
| 改动 | 为什么有效 | 如何约束结果变化 |
| --- | --- | --- |
| 原生RGB排布、类别筛选、C++ NMS | 减少Python循环和小对象 | 与参考数组/OpenCV NMS对照 |
| COCO GPU类别归约 | 在GPU取每候选最大类别分数，84通道→6通道 | 候选数不变，传输量减少92.86% |
| 合并两路NMS工作 | 一次输出同时供跟踪0.05和评估0.001使用 | 保留各自阈值、相同IoU与数量上限 |
| CUDA Graph | 重放固定形状H2D/TRT/D2H，减少逐次提交开销 | 同形状对照日志逐字节一致 |
| 缓冲复用、有界队列 | 减少重复分配，避免积压所有帧 | 队列满时等待，不丢帧 |

![Graph和网页调度](插图/04_Graph与网页调度.png)

640×384固定流水线的Graph对照为88.90→104.14 FPS，约+17.1%；512×320约99→110.49 FPS。1280没有稳定收益，当前只在矩形速度方案默认开启Graph。

GPU归约做了36组物理GPU数组验证并检查真实视频样本。原生后处理阶段的标准640完整429帧，三份检测/轨迹日志与参考一致。后续C格式化JSON使用 `%.17g` 并释放GIL：JSON文本写法可能不同，解析后的浮点数逐项一致，轨迹CSV仍逐字节一致。

没有收益的尝试也保留结论：OpenCV改单线程、某些提前提交/重叠推理方式，没有稳定超过最终实现；保留2线程与简单有界流水线。小模型的速度常受提交、CPU和同步限制，添加更多线程或更复杂流水线不一定更快。
''')

page('06　真实网页为何会比独立脚本慢', '''
独立流水线接近或超过百帧后，首次真实网页测试仍只有约70 FPS。网页要生成JPEG预览、序列化大量检测记录、响应状态请求；部分Python代码持有GIL时，其他线程拿不到执行机会，GPU就可能等下一帧。

最终联合调整：检测JSON在C中格式化并释放GIL；Python线程切换间隔设为1 ms；速度方案预览改为640宽、4次/秒，标准预览仍是1280宽、8次/秒。预览只抽取展示，不减少检测或保存帧数。

| 同一真实网页任务 | 调度前 FPS | 调度后 FPS | 变化 |
| --- | ---: | ---: | ---: |
| MOT20-01 / 640标准 | 41.27 | 56.00 | +35.7% |
| MOT20-01 / 640×384固定 | 71.76 | 104.17 | +45.2% |
| MOT20-01 / 512×320固定 | 69.68 | 116.75 | +67.5% |
| MOT17-09 / 512×320固定 | 86.99 | 121.74 | +40.0% |

这三项是联合改进，没有独立消融，不能把整份收益全部归给JSON或某一个参数。图4右侧显示对应网页对照。

最新标准640平均推理与传输约10.96 ms、跟踪17.08 ms、输出线程绘图日志6.85 ms，处理吞吐56 FPS。推理已不是最大单段；单纯换更快的检测引擎，整体仍会受GMC与关联限制。

NVENC是专用视频编码单元，不等于GPU神经网络算力；使用它可减少CPU编码负担。当前仍有CPU解码和预处理，未声称整条链路“零拷贝”，也未测摄像头端到端延迟。达到离线100 FPS说明处理能力超过该段源片25 FPS，不等于摄像机从曝光到显示只有10 ms。

当前40W与最高频率仅本次开机有效。重启后可双击“启用高性能模式”，由系统授权窗口完成；启动网页本身不要求管理员密码。持续运行要结合温度、功耗和真实场景测试，当前没有以关闭预览或删除保存来获得表中的速度。
''')

profilerows=[]
for k,label in zip(order,labels):
    s=web[k]['state'];profilerows.append([label,f"{fps(k):.2f}",f"{s['elapsed']:.2f}",f"{429/s['elapsed']:.2f}",f"{ap(k):.2f}%",f"{ap(k,'map50_95'):.2f}%"])
page('07　五种正式方案：按场景选，不混算收益', '''
![速度与精度](插图/05_速度精度取舍.png)

'''+table(['MOT20-01 / 429帧','处理FPS','总秒','整任务FPS','AP50','AP50:95'],profilerows)+'''

**1280小目标**是网页COCO默认选项，保留完整GMC；**640标准**保留相同GMC；**640×384移动**把GMC降采样由2改4、角点1000改300；**640×384固定**关闭GMC；**512×320极速**进一步缩小输入。三个速度选项同时开启Graph、紧凑绘图和较低频预览。

矩形输入减少横向画面的无效填充；输入面积640×384比640×640少40%，512×320少60%。实际letterbox像素、网络边界与输出网格不同，所以不能保证检测相同。640×384固定比标准640快86.0%，AP50低0.79个百分点；极速快108.5%，AP50低4.16个百分点。不能把它们称为同精度加速。

固定机位优先试640×384；移动镜头用带轻量GMC方案；远处小目标、拥挤遮挡选1280，再按需求考虑微调或切片。UAV640另测36.49 FPS，VisDrone1280另测14.73 FPS，说明百帧结论不会自动覆盖所有视频和模式。
''')

introws=[]
for vid,label in [('MOT20-01-raw.mp4','MOT20-01'),('VisDrone-val-uav0000137_00458_v-raw.mp4','VisDrone')]:
    rs=[next(x for x in q if x['video']==vid and x['engine']==f'coco_yolov12n_640_{pr}.engine' and x['confidence']==.001) for pr in ['fp16','int8']]
    introws.append([label,f"{rs[0]['metrics']['map50']*100:.2f}%",f"{rs[1]['metrics']['map50']*100:.2f}%",f"{(rs[1]['metrics']['map50']-rs[0]['metrics']['map50'])*100:.2f}个百分点"])
page('08　为什么画面看着准，mAP却不高', '''
排查发现两项实际原因。第一，早期把跟踪阈值0.05也用于AP，提前删掉了低分候选；现在评估保留到0.001。同样固定60帧的FP32，MOT AP50从47.84%到56.85%，VisDrone从30.80%到34.39%。这是**评估口径修正**，不能称为模型学到了更多目标。

第二，640输入对航拍小目标信息不足。VisDrone原片2688×1512，缩到640时，GT目标短边中位数约14.05像素，9.34%短边不足8像素；1280时中位数28.10像素，不足8像素比例降到0.13%。

![同帧人工GT与不同输入检测](插图/08_小目标同帧对照.png)

上图显示同一帧、同一区域：人工GT为绿框，640/1280检测为蓝框，图示显示阈值统一0.30；实际评估仍用0.001。框更多不自动等于更准。

平均数也会与观感不同：固定VisDrone帧640 FP16的car AP50:95约44.29%，person17.56%，bicycle1.38%，motorcycle4.25%，四类等权平均只有16.87%。人眼容易注意近处清楚的大目标，AP还会处罚漏检、错类与边界偏差。

已核对BGR/RGB、NCHW、/255、填充114、逆缩放、类别表、视频与GT同步、模型哈希；模型输出没有独立objectness，不能额外相乘。同640口径的FP32与原TRT FP16 AP50差异小于0.14个百分点，未发现足以解释低分的FP16精度破坏。当前YOLOv12n及COCO权重对航拍的适应性仍有限；硬件等级提高的是计算能力，不会自动改变训练数据和模型能力。

自有视频无人工GT时只能检测，mAP显示不适用；错误视频/尺寸/帧数/类别标注也会拒评。不能拿置信度或已画框的视频替代人工真值。
''')

page('09　INT8和“NPU”：实测后为何仍选GPU FP16', '''
![DLA与INT8隔离延迟](插图/06_DLA与INT8.png)

**INT8。** 使用EntropyCalibration2，5段非测试视频各32帧、共160真实帧校准，保留校准缓存和帧号；不是随机设动态范围。模型仍有FP16/FP32回退和69个Reformat层。交替3轮测试FP16平均8.92 ms、INT8 10.36 ms，延迟增加约16.0%。固定60个留出帧、相同0.001阈值的AP如下，不能与全429帧表直接当成同一成绩。

'''+table(['固定60帧 / 640','FP16 AP50','INT8 AP50','变化'],introws)+'''

本版PTQ同时损失速度与精度，保留实验入口但不推荐。格式转换、敏感浮点层与回退都可能抵消低精度算术收益。显式Q/DQ、敏感层保留FP16或QAT值得另测，本次没有这些方案的收益数字。

**DLA。** Orin上可尝试的独立深度学习加速器是NVDLA，TensorRT查询到2个核心。本次实际构建和测量的是**core 0**，未测双核并行。纯DLA在 `/model.2/Shape` 失败；允许GPU回退后成功，但分成16个DLA子图、58个Reformat，还有GPU矩阵与其他算子。隔离GPU前后交叉均值约10.33 ms，DLA/GPU混合39.68 ms；混合完整MOT20-01只有20.48 FPS。

这说明当前YOLOv12的图结构和分割代价不适合直接卸载，不能把GPU和两个DLA的理论TOPS直接相加。PVA面向特定视觉任务，也不是另一个可无修改运行YOLO的通用NPU；本次未实现PVA卸载。已有效利用CPU并行、GPU TensorRT/CUDA Graph和NVENC。若要DLA有价值，应另选支持算子更友好的模型，或让它处理独立轻量ReID任务后重新验证，不能承诺必然提速。
''')

trackrows=[]
for label,r in zip(['640标准','640×384轻量GMC','640×384固定','512×320固定'],t['runs']):
    m=r['metrics'];trackrows.append([label,f"{m['idf1']*100:.2f}%",int(m['num_switches']),int(m['num_misses'])])
page('10　BoT-SORT、GMC和ReID分别带来什么', '''
检测只找每帧的框；BoT-SORT用运动预测、IoU/分数关联、低分框补关联和轨迹管理，给目标连续ID。GMC补偿相机运动；ReID再通过另一张网络比较目标外观。这些改善的是身份关联能力，本身会增加计算，不能把mAP变化归功于它们。

当前代码严格限制同类别关联，空帧也更新轨迹；默认关闭ReID。辅助身份评估使用MOT20-01官方ID、IoU≥.5、全部可见度，按规则忽略非计分区域，使用motmetrics1.4；不是官方MOTChallenge评测。

'''+table(['同429帧辅助对照','IDF1','身份切换','漏检数'],trackrows)+'''

极速模式虽然身份切换更少，但也漏掉更多目标、IDF1下降。只数ID切换或唯一ID减少，不能证明跟踪变好。目前没有同GT下“另一跟踪器→BoT-SORT”的成对质量提升百分比。

![ReID阶段代价](插图/07_ReID代价.png)

配套ReID是ResNeSt-50x + WITH_NL、2048维，实际用PyTorch FP32。FP16曾产生非有限特征，因此回退FP32。固定同一检测缓存，UAV100帧跟踪阶段19.06→996.91 ms（52.3倍）；MOT20-01为22.57→310.17 ms，MOT20-08为19.46→245.13 ms。无人机100帧完整链路另测0.956 FPS。

MOT实验使用配套无人机领域ReID权重，没有验证过行人专用权重，也未用GT证明开启它后的IDF1收益。当前默认关闭是有实测依据的选择；轻量领域匹配ReID、TensorRT部署、仅在遮挡/歧义时调用，是后续可研究的方向，收益需重测。
''')

page('11　代码、运行方式和结果应该去哪里找', '''
日常双击 `Jetson_视频工作台/双击启动网页`，选择UAV或COCO、原始视频与方案，点击启动，完成后回放。自己的MP4也能用；要评mAP再提供同名 `.gt.json`。默认COCO为1280小目标模式；固定横向画面可试640×384。转发请发送整个独立目录或“可转发.zip”，在兼容Jetson上解压后双击。

| 独立包内路径 | 职责 |
| --- | --- |
| app/web/、app/server.py、jobs.py | 网页、播放器、接口、单任务管理与取消 |
| app/jetson_runtime.py、pipeline_workers.py | 有界流水线、指标、解码、绘图、NVENC与预览 |
| runtime/backend.py、frame_ops.py | TensorRT/CUDA、预后处理、原生库与回退 |
| runtime/native/ | C / C++ / CUDA源码与ARM64共享库 |
| vendor/botsort/tracker/ | 多类别BoT-SORT、GMC、轨迹状态管理 |
| app/evaluation.py、fast_coco.py | GT校验、COCO AP与原生匹配适配 |
| models/、environment/、browser/ | ONNX/引擎及元数据、离线wheel、ARM64浏览器 |
| tools/、tests/、docs/ | 构建导出、复现实验、验证和代码导航 |
| videos/、outputs/、reports/ | 原始素材、逐次结果、报告和实验依据 |

运行不依赖旧工程、PyTorch或FastReID。全新兼容Jetson仍需有JetPack的CUDA/TensorRT/OpenCV及系统venv/tk/ffmpeg；不是Windows双击HTML就能运行的网页。版本不匹配时用包内ONNX在目标机重建FP16，引擎构建要等待数分钟。

![历史ReID同帧关联差异](插图/09_ReID交汇放大.png)

上图为源帧67同检测缓存的交汇区域放大，左右ReID开关改变了ID分配，不能凭这张图判定哪边身份正确。直观视频仍在工作区 `视频结果/直观对比/`；最新六种结果集中在 `视频结果/百帧优化_20261015/`。完整视频按源FPS播放，比较速度应看本报告的处理耗时。
''')

released=cleanup['estimated_released_bytes']/1e9
page('12　验证、文件整理与下一步优化', f'''
**已完成验证。** 29项单元测试通过；36组物理GPU归约验证及真实视频样本检查通过；标准640完整429帧的原生优化结果等价；Graph开关日志一致；C JSON解析数值一致。真实网页跑完8个任务，检查GPU/NVENC、全部帧数与原尺寸、H.264回放、取消Graph、拒绝并发、错误GT不造分数。六个单列视频共2670帧完整解码通过。

上轮可转发包还在中文带空格的新路径全新解压、离线安装并一键启动，完整429帧约112.57 FPS。该次是历史验收值；更新报告后的包另做全新解压、离线启动、完整429帧GPU任务与H.264解码，以及SHA/CRC/清单验证。新包验收值和对应哈希见交付包旁JSON，历史测试按旧哈希单独保留。

**本次目录整理。** 主程序保持原位；早期部署、训练代码、权重和原始移交文件作为一组移到 `开发归档/早期部署工程/`，保留组内相对路径；旧网页、数据准备和精度排查各归入开发归档。工作区根只保留主程序、交付ZIP、报告入口、原始视频、结果视频、开发归档和导航。

删除已验证MP4完好的旧MJPEG AVI、与已解压全部文件逐项哈希相同的下载ZIP、可再生字节码；大文件重复副本经SHA256校验后用硬链接共享存储，路径仍保留。**本轮删除与去重估计释放{released:.2f} GB**，其中直接删除{cleanup['deleted_bytes']/1e9:.2f} GB、去重{cleanup['dedup_released_bytes']/1e9:.2f} GB；文件系统可用空间实测净变化{cleanup['filesystem_free_delta_bytes']/1e9:.2f} GB，删除/去重估计使用逻辑文件大小，实际净变化还受文件分配块和同期系统写入影响。之前已清理的约3.10 GB单独记录，没有重复计入本轮。

原始视频、GT、独有JPEG、模型与训练权重、正式结果、历史实测日志、当前虚拟环境与浏览器状态均保留。硬链接输入请视为只读；要编辑时另存副本。逐文件原路径、新路径、删除原因、SHA和字节数见 `../validation/目录整理_20261015.json`。

**下一步按收益与风险排序。** ①精度优先：目标领域微调、小目标切片或更大模型，先用独立GT评AP；②移动镜头速度：进一步检查GMC与关联CPU耗时、硬件解码/预处理是否值得，保持逐帧验证；③身份优先：轻量、领域匹配ReID按需调用，测IDF1和总耗时；④量化/DLA：重做显式量化或更兼容的模型结构，必须同时看速度与质量。以上尚未实测，不承诺具体提升，也没有证据说设备已达到所有工作负载的绝对上限。

**复现入口。** `tools/benchmark_pipeline.py`、`benchmark_ceiling.py`、`tests/browser_ceiling.cjs`、`compare_tracking.py` 提供实验代码；构建见 `tools/build_engine.py`、`export_shape.py` 及 `runtime/native/`。本报告重生成：在包内执行 `python3 tools/make_technical_report.py`（需开发绘图库）。数据索引见下一页，读者可核对每个结论的来源。
''')

page('附录　数据索引、复现边界与计算公式', '''
| 证据文件（相对于包根） | 对应内容 |
| --- | --- |
| reports/技术报告/数据/历史后端与ReID原始摘要.json | 原PyTorch/TRT及三组固定检测ReID数据，含原文件SHA |
| reports/validation/performance_comparison.json | 相同视频/尺寸的流水线前后、阶段时间、日志一致性 |
| reports/validation/ceiling_comparison.json | Graph/DLA探索、网页前后、全部帧最新速度与AP |
| reports/benchmarks/paired_benchmark.json | 两段固定60帧，门槛/输入/INT8的AP，冻结帧号 |
| reports/benchmarks/alternating_latency.json | FP16/INT8三轮交替隔离延迟 |
| reports/validation/dla_audit.json、dla_layers.json | DLA核心查询、构建失败/回退、层类型与测时 |
| reports/audit_torch_640.json、audit_trt_640.json | 原FP32/FP16同口径审计 |
| reports/validation/tracking_ceiling.json | 辅助身份质量及忽略规则、GT与轨迹SHA |
| reports/validation/final_numeric_equivalence.json | C JSON数值等价，轨迹CSV一致 |
| reports/validation/browser_ceiling.json | 真实网页8个完整任务与回放/取消/并发验收 |
| reports/validation/目录整理_20261015.json | 本轮移动、删除和去重逐文件清单 |
| reports/技术报告/数据/报告来源校验.json | 本报告读入数据与插图SHA、生成器SHA |

原始大体积实验日志以tar.gz保存在本机 `reports/benchmarks/`，转发ZIP提供摘要和验证代码，不含这些大归档、历史输出或本地虚拟环境。报告中的旧绝对路径用于追溯历史，旧实验输出已归档时不要把该路径当成现存文件。历史开发工程已整体移动，重新导出请把 `--project-root` 指向 `开发归档/早期部署工程`。

吞吐提升＝新FPS÷旧FPS−1；加速倍数＝新FPS÷旧FPS；耗时降低＝1−新秒数÷旧秒数。AP变化用**百分点**，如57.12%→56.33%是低0.79个百分点。检测推理倍数、完整流水线倍数、不同输入尺寸的变化不能直接相乘成总收益。

当前评测范围是已列视频、对应类别和现有权重。没有全COCO/全VisDrone测试集成绩，没有摄像头延迟/断线恢复测试，没有双DLA并行、ReID TensorRT或QAT收益，最新版本也未重新跑30分钟长稳。早期另有30分钟不编码视频循环，不能替代最新百帧网页的长稳结论。

报告与源码、测试和数据一同交付；后续换模型、帧率、分辨率或目标密度，应重新测量。先保持计分规则与输入一致，再比较提速或质量，才能判断改动是否真的有用。
''')

revision=ROOT/'reports/validation/small_input_corrected.json'
if revision.is_file():
    updated=read('reports/validation/small_input_corrected.json')
    values=[x for x in updated if x['pipeline']['frames']==3315]
    extra=[]
    for label,x in zip(['原512','512去重优化','512分区增强'],values):
        z=x['pipeline'];e=z['evaluation']
        extra.append([label,f"{e['map50']*100:.2f}%",f"{e['map50_95']*100:.2f}%",f"{z['processing_fps']:.2f}"])
    sections=pages[0].split('\n\n',1)
    pages[0]=sections[0]+'\n\n> 后续512精度更新见“更新补充”页。前文512速度/精度保留最初NMS0.7基准；当前网页极速改用NMS0.55，并新增每帧五次推理的分区增强选项。\n\n'+sections[1]
    page('更新补充　512密集人群精度与编码修复', '''
针对用户反馈的MOT20密集人群，使用同一个512×320 FP16引擎完成了MOT20-05全部3315帧的新对照，评估GT、门槛和COCO规则保持一致。没有跳帧。

'''+table(['MOT20-05 / 3315帧','AP50','AP50:95','处理FPS'],extra)+'''

“极速512”将NMS IoU从0.7调到0.55，减少错位重复框，不增加推理次数；不能无限降低阈值，以免抑制相邻真实行人。另一选项“人群细节512”每帧全图＋四个重叠区域，共五次推理，恢复原图坐标、过滤内部截断框并按类合并，BoT-SORT每源帧只更新一次。增强选项降低新轨迹门槛以便显示低分小目标；其IDF1尚未测量，AP变化来自检测，不来自跟踪门槛。

原单图、去重优化与分区增强的实际结果视频已单列在工作区 `视频结果/512密集人群优化_20261015/`。完整图表、实际同帧截图、冻结40帧参数探索及MOT20-02/VisDrone交叉检查见 [512密集人群精度优化报告](../512密集人群优化/512密集人群精度优化报告.html)。锐化、亮度、CLAHE和插值没有可靠收益，未纳入默认链路。

还修复826像素等宽度的BGR行与OpenCV/GStreamer四字节对齐问题。先补齐输入行，再在BGRx阶段裁回原宽度，NVENC仍保存原画面尺寸。修复前的花屏/黑屏视频没有作为最终结果；表中三种方案都已重跑，完整3315帧解码与实际图像检查通过。旧错误不影响从原帧计算的检测AP。

处理FPS排除加载预热和最终评估。原版用同生产流水线基准脚本，新方案用正常网页接口；后台负载也在变化，不把两者速度差归为单项收益。极速方案保留高吞吐，五视图增强明确付出多次推理代价，不能声称百帧。原640/1280方案保留。

最新验证为35项单元测试，真实浏览器完成两方案429帧、NVENC、H.264回放、1/5次检测计数、分区中途取消和并发拒绝。新数据见 `reports/validation/small_input_corrected.json` 和 `browser_small_input.json`，旧百帧数据仍按历史版本保留。
''')
    latest=pages.pop();pages.insert(-1,latest)
md = '# Jetson 视频工作台：部署、提速与方案对比\n\nYOLOv12 · TensorRT · BoT-SORT · ReID · INT8 · DLA\n\n'+ '\n\n'.join(pages)
(OUT/(TITLE+'.md')).write_text(md)
css = '''@page{size:A4;margin:13mm 14mm 14mm}*{box-sizing:border-box}body{margin:0;color:#203344;font:11.3pt/1.64 'Noto Sans CJK JP',sans-serif}article{break-after:page}article:last-child{break-after:auto}.strap{font-size:8pt;color:#748695;border-bottom:1px solid #d9e4e9;padding-bottom:6px;margin-bottom:15px;letter-spacing:1px}h2{font-size:20pt;line-height:1.35;color:#087c74;margin:0 0 15px}p{margin:10px 0}strong{color:#126b68}img{display:block;max-width:100%;max-height:260px;object-fit:contain;margin:10px auto}table{border-collapse:collapse;width:100%;font-size:9.2pt;line-height:1.5;margin:12px 0;break-inside:avoid}th{background:#e9f5f2;color:#12685f}td,th{padding:6px;border:1px solid #cfdee3;text-align:left}code{font:8.5pt/1.45 monospace;background:#edf3f5;overflow-wrap:anywhere}a{color:#167e9c}pre{white-space:pre-wrap;font-size:9pt}h2,th{break-after:avoid}article p:has(img){break-inside:avoid}@media screen{body{background:#e8eef1}article{background:white;width:210mm;min-height:280mm;margin:22px auto;padding:17mm;box-shadow:0 4px 25px #bbccd4}}'''
html='<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>Jetson 部署与提速技术报告</title><style>'+css+'</style><body>'
for n,part in enumerate(pages,1):
    html+=f'<article><div class="strap">JETSON 视频工作台　/　技术报告　/　2026.10.15　/　{n:02d}</div>'+markdown.markdown(part,extensions=['tables','fenced_code'])+'</article>'
html+='</body></html>'
(OUT/(TITLE+'.html')).write_text(html)
subprocess.run([str(ROOT/'browser/chrome-linux-arm64/chrome'),'--headless','--no-sandbox','--disable-gpu',
    '--no-pdf-header-footer',f'--print-to-pdf={OUT/(TITLE+".pdf")}',(OUT/(TITLE+'.html')).as_uri()],
    check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)

# An editable companion, using the same source paragraphs, charts and tables.
try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();style=doc.styles['Normal'];style.font.name='Noto Sans CJK JP';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK JP')
    section=doc.sections[0];section.top_margin=Inches(.55);section.bottom_margin=Inches(.55)
    for n,part in enumerate(pages):
        if n:doc.add_page_break()
        lines=part.splitlines();i=0
        while i<len(lines):
            line=lines[i].strip();i+=1
            if not line:continue
            if line.startswith('## '):doc.add_heading(line[3:],1)
            elif line.startswith('!['):doc.add_picture(str(OUT/line.split('](',1)[1].rstrip(')')),width=Inches(6.2))
            elif line.startswith('|'):
                rows=[line]
                while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
                values=[[x.strip() for x in row.strip('|').split('|')] for j,row in enumerate(rows) if j!=1]
                tab=doc.add_table(rows=0,cols=len(values[0]));tab.style='Light Shading Accent 1'
                for row in values:
                    cells=tab.add_row().cells
                    for cell,val in zip(cells,row):cell.text=val.replace('`','')
            else:doc.add_paragraph(line.replace('**','').replace('`',''))
    doc.save(OUT/(TITLE+'.docx'))
except ImportError:
    pass
catalog.append(dict(path=str(snapshot.relative_to(ROOT)),bytes=snapshot.stat().st_size,sha256=sha(snapshot)))
source=Path(__file__).resolve()
manifest=dict(report_date='2026-10-15',new_inference=False,generator=dict(path=str(source.relative_to(ROOT)),sha256=sha(source)),
              sources=catalog,figures=[dict(path=str(f.relative_to(OUT)),sha256=sha(f)) for f in sorted(FIG.glob('*.png'))],
              outputs=[dict(path=f.name,bytes=f.stat().st_size,sha256=sha(f)) for f in sorted(OUT.glob(TITLE+'.*'))])
(DATA/'报告来源校验.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
(OUT/'README_阅读说明.md').write_text('''# 技术报告\n\n优先打开PDF；HTML适合浏览器阅读，Markdown便于核对原文，DOCX可编辑。\n\n正文、更新补充与证据附录，覆盖部署、TensorRT FP16、串行瓶颈、并行/NVENC、原生后处理、GPU归约、Graph、网页调度、精度、INT8/DLA、BoT-SORT/ReID、使用与目录整理。512密集人群最新精度结果另有更新补充页和同级独立报告。插图来自实测JSON或已有同帧可视化。\n\n数据与生成代码一同保留。重生成会复用历史测量，不运行新推理。\n''')
print(OUT/(TITLE+'.pdf'))
